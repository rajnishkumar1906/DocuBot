# app.py
import os
from pathlib import Path
from typing import List, Dict
import streamlit as st
import PyPDF2
import docx
import re
from dotenv import load_dotenv
import google.generativeai as genai

from langchain_huggingface import HuggingFaceEmbeddings
from langchain_chroma import Chroma
from langchain.schema import Document

# ----------------------------
# Load Gemini API Key
# ----------------------------
load_dotenv()
genai.configure(api_key=os.getenv("GEMINI_API_KEY"))

# ----------------------------
# Helpers
# ----------------------------
def extract_text(file):
    try:
        if file.type == "text/plain":
            return file.read().decode("utf-8")
        elif file.type == "application/pdf":
            reader = PyPDF2.PdfReader(file)
            return "\n".join([page.extract_text() or "" for page in reader.pages])
        elif file.type in ["application/vnd.openxmlformats-officedocument.wordprocessingml.document","application/msword"]:
            doc = docx.Document(file)
            return "\n".join([p.text for p in doc.paragraphs])
    except Exception:
        return ""

def chunk_text(text, chunk_size=500, overlap=50):
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(Document(page_content=text[start:end], metadata={"source": "uploaded_file"}))
        start += chunk_size - overlap
    return chunks

def get_embedder():
    return HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2", model_kwargs={"device":"cpu"})

def get_vectorstore(documents: List[Document]):
    embeddings = get_embedder()
    vs = Chroma(collection_name="docubot", embedding_function=embeddings)
    vs.add_documents(documents)
    return vs

def retrieve(vs: Chroma, query: str, k: int = 4):
    return vs.as_retriever(search_kwargs={"k": k}).invoke(query)

def synthesize_answer(query: str, docs: List[Document], chat_history: List[tuple]) -> str:
    history_str = "\n".join([f"{role}: {msg}" for role, msg in chat_history])
    context = "\n\n".join([d.page_content for d in docs])
    prompt = f"""
You are DocuBot, an AI assistant. Answer questions using only the uploaded documents.

Conversation History:
{history_str}

Context:
{context}

Question:
{query}

Answer:
Provide a concise answer based ONLY on the context. If unknown, say you don't have enough information.
"""
    model = genai.GenerativeModel("gemini-1.5-pro")
    try:
        response = model.generate_content(prompt)
        return response.text.strip()
    except Exception as e:
        return f"Error generating answer: {e}"

def answer_query(vs: Chroma, query: str, chat_history: List[tuple], file_uploaded: bool, k: int = 4) -> Dict[str, any]:
    greetings = ["hi", "hello", "hey", "good morning", "good evening", "good afternoon"]
    thanks_phrases = ["thanks", "thank you", "nice", "cool", "great"]
    meta_questions = ["discuss", "about the file", "about the pdf", "about my file", "about the document", "about the documents"]

    q = query.lower().strip()
    if q in greetings:
        return {"answer": "👋 Hello! I'm DocuBot.", "sources": []}
    if any(p in q for p in thanks_phrases):
        return {"answer": "You're welcome!", "sources": []}
    if "how are you" in q:
        return {"answer": "😅 I’m just code, but ready to assist!", "sources": []}
    if any(p in q for p in meta_questions):
        return {"answer": "Please upload a document first." if not file_uploaded else "Yes, I have processed your files.", "sources": []}

    try:
        docs = retrieve(vs, query, k=k)
        if not docs:
            return {"answer": "No relevant content found.", "sources": []}
    except Exception:
        return {"answer": "Error retrieving documents.", "sources": []}

    answer = synthesize_answer(query, docs, chat_history)
    sources = list({d.metadata.get("source", "unknown") for d in docs})
    return {"answer": answer, "sources": sources}

# ----------------------------
# Streamlit App
# ----------------------------
st.set_page_config(page_title="DocuBot", layout="wide")
st.title("🤖 DocuBot RAG Chat")

if "vs" not in st.session_state:
    st.session_state.vs = None
if "history" not in st.session_state:
    st.session_state.history = []
if "file_uploaded" not in st.session_state:
    st.session_state.file_uploaded = False

# ----------------------------
# File Upload
# ----------------------------
uploaded_files = st.file_uploader("📂 Upload TXT, PDF, DOCX", type=["txt","pdf","docx"], accept_multiple_files=True)

if uploaded_files:
    all_docs = []
    for file in uploaded_files:
        text = extract_text(file)
        if text.strip():
            chunks = chunk_text(text)
            all_docs.extend(chunks)
    st.session_state.vs = get_vectorstore(all_docs)
    st.session_state.file_uploaded = True
    st.success(f"Indexed {len(all_docs)} chunks from uploaded documents!")

# ----------------------------
# Chat Input
# ----------------------------
query = st.text_input("Ask me anything about your documents:", key="user_prompt")

if query and st.session_state.vs:
    st.session_state.history.append(("User", query))
    result = answer_query(st.session_state.vs, query, st.session_state.history, st.session_state.file_uploaded)
    st.session_state.history.append(("DocuBot", result["answer"]))

# ----------------------------
# Display full conversation
# ----------------------------
st.markdown("## 💬 Conversation")
for role, msg in st.session_state.history:
    if role == "User":
        st.markdown(f'<div style="background: #0072ff; color: white; padding: 10px; border-radius: 10px; margin-bottom:5px;">{msg}</div>', unsafe_allow_html=True)
    else:
        st.markdown(f'<div style="background: #f0f0f0; color: #333; padding: 10px; border-radius: 10px; margin-bottom:5px;">{msg}</div>', unsafe_allow_html=True)
